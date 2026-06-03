# -*- coding: utf-8 -*-
"""
reporte.py — Genera el reporte QA Checklist (con evidencias incrustadas y Dashboard).
Uso: python reporte.py <reporte_data.json> <salida.xlsx>
"""
import sys
import os
import json
import struct
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import PieChart, BarChart, Reference

# ── Paleta (igual al template de referencia) ──────────────────
NAVY   = "001F5B"
COBALT = "0047AB"
GRAYBL = "E9EDF3"
BANDED = "F5F7FA"
WHITE  = "FFFFFF"
TXT    = "1A2332"
TXTHDR = "37474F"

EST_FILL = {
    "Aprobado": "C8E6C9", "Fallido": "FFCDD2", "Pendiente": "FFE0B2",
    "En Progreso": "BBDEFB", "Bloqueado": "E1BEE7",
}
PRI_FILL = {"Alta": "FFEBEE", "Media": "FFF8E1", "Baja": "E8F5E9"}

CHECK = "✅ Checklist"
DASH  = "📊 Dashboard"

thin = Side(style="thin", color="D5DBE3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def solid(color):
    return PatternFill("solid", fgColor=color)


def png_size(path):
    """Lee ancho/alto del PNG sin depender de nada externo."""
    try:
        with open(path, "rb") as f:
            head = f.read(24)
        w, h = struct.unpack(">II", head[16:24])
        return w, h
    except Exception:
        return 1600, 900


def construir_checklist(wb, meta, filas):
    ws = wb.active
    ws.title = CHECK
    ws.sheet_view.showGridLines = False

    # Anchos de columna (M = Evidencia ampliada)
    anchos = [12.6, 26, 15, 18, 35, 37, 42, 35, 14, 16.7, 17.7, 14, 42, 30]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Fila 1: título ──
    ws.merge_cells("A1:N1")
    c = ws["A1"]
    c.value = "EVEREST TEST  |  SISTEMA DE GESTIÓN DE CALIDAD  |  QA CHECKLIST"
    c.fill = solid(NAVY)
    c.font = Font(color=WHITE, bold=True, size=16)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 49.5

    # ── Fila 2: bloque encabezado ──
    pares = [
        ("A", "Módulo:", "B", meta.get("moduloHeader", "HISTORIAS CLINICAS")),
        ("C", "Versión:", "D", meta.get("version", "1.0.0")),
        ("E", "Ambiente:", "F", meta.get("ambiente", "PRUEBAS")),
        ("G", "Fecha Inicio:", "H", meta.get("fecha", "")),
        ("I", "Fecha Fin:", "J", meta.get("fecha", "")),
        ("K", "Responsable Dev:", "L", meta.get("desarrollador", "")),
    ]
    for lc, lbl, vc, val in pares:
        a = ws[f"{lc}2"]; a.value = lbl
        a.fill = solid(COBALT); a.font = Font(color=WHITE, bold=True, size=10)
        a.alignment = Alignment(horizontal="center", vertical="center")
        b = ws[f"{vc}2"]; b.value = val
        b.fill = solid(GRAYBL); b.font = Font(color=TXTHDR, bold=True, size=10)
        b.alignment = Alignment(horizontal="center", vertical="center")
    # M2:N2 relleno encabezado
    for col in ("M", "N"):
        ws[f"{col}2"].fill = solid(GRAYBL)
    ws.row_dimensions[2].height = 24

    # ── Fila 3: espaciador ──
    ws.row_dimensions[3].height = 6

    # ── Fila 4: headers de columna ──
    headers = [
        "ID Caso\nPrueba", "Cliente", "Ambiente", "Módulo", "Flujo /\nFuncionalidad",
        "Escenario /\nCaso de Prueba", "Resultado\nEsperado", "Resultado\nObtenido",
        "Estado", "Prioridad", "Tester Dev", "Fecha\nEjecución", "Evidencia", "Observaciones",
    ]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(4, i, h)
        cell.fill = solid(COBALT)
        cell.font = Font(color=WHITE, bold=True, size=9)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[4].height = 31.5

    # ── Filas de datos ──
    fila0 = 5
    for idx, r in enumerate(filas):
        fila = fila0 + idx
        banda = BANDED if idx % 2 else WHITE
        valores = [
            r.get("id", ""), r.get("cliente", ""), r.get("ambiente", "PRUEBAS"),
            r.get("modulo", "HISTORIAS CLINICAS"), r.get("flujo", ""), r.get("escenario", ""),
            r.get("esperado", ""), r.get("obtenido", ""), r.get("estado", ""),
            r.get("prioridad", ""), r.get("tester", "QA Automatization"),
            r.get("fecha", ""), "", r.get("observaciones", ""),
        ]
        for ci, val in enumerate(valores, start=1):
            cell = ws.cell(fila, ci, val)
            cell.fill = solid(banda)
            cell.font = Font(color=TXT, size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        # Evidencia incrustada (col M = 13)
        ev = r.get("evidencia", "")
        alto_fila = 70
        if ev and os.path.exists(ev):
            w, h = png_size(ev)
            ancho_px = 250
            alto_px = round(ancho_px * h / w)
            img = XLImage(ev)
            img.width = ancho_px
            img.height = alto_px
            ws.add_image(img, f"M{fila}")
            alto_fila = max(alto_fila, alto_px * 0.75 + 6)
        ws.row_dimensions[fila].height = alto_fila

    ult = fila0 + len(filas) - 1  # última fila de datos

    # ── Formato condicional Estado (I) y Prioridad (J) ──
    rango_est = f"I{fila0}:I{ult}"
    for estado, color in EST_FILL.items():
        ws.conditional_formatting.add(
            rango_est,
            FormulaRule(formula=[f'$I{fila0}="{estado}"'], fill=solid(color), stopIfTrue=False),
        )
    rango_pri = f"J{fila0}:J{ult}"
    for pri, color in PRI_FILL.items():
        ws.conditional_formatting.add(
            rango_pri,
            FormulaRule(formula=[f'$J{fila0}="{pri}"'], fill=solid(color), stopIfTrue=False),
        )

    # ── Dropdowns ──
    dv_estado = DataValidation(type="list", formula1='"Pendiente,En Progreso,Aprobado,Fallido,Bloqueado"', allow_blank=True)
    dv_pri = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
    dv_mod = DataValidation(type="list", formula1='"HISTORIAS CLINICAS,AGENDAMIENTO,FACTURACIÓN,LABORATORIO,AUTENTICACIÓN,CONSENTIMIENTOS,ORDENAMIENTOS,REIMPRESION,NOTA"', allow_blank=True)
    ws.add_data_validation(dv_estado); ws.add_data_validation(dv_pri); ws.add_data_validation(dv_mod)
    dv_estado.add(f"I{fila0}:I498"); dv_pri.add(f"J{fila0}:J498"); dv_mod.add(f"D{fila0}:D498")

    # ── Totales ──
    ft = ult + 1
    ws.merge_cells(f"A{ft}:H{ft}")
    a = ws[f"A{ft}"]; a.value = "TOTALES AUTOMÁTICOS →"
    a.fill = solid(NAVY); a.font = Font(color=WHITE, bold=True, size=9)
    a.alignment = Alignment(horizontal="center", vertical="center")
    et = ["Total Casos", "Aprobados", "Fallidos", "Pendientes", "En Progreso", "Bloqueados"]
    for i, txt in enumerate(et):
        col = 9 + i
        cell = ws.cell(ft, col, txt)
        cell.fill = solid(COBALT); cell.font = Font(color=WHITE, bold=True, size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[ft].height = 18

    fv = ft + 1
    ws.merge_cells(f"A{fv}:H{fv}")
    a = ws[f"A{fv}"]; a.value = "Ver Dashboard → hoja 📊 Dashboard"
    a.fill = solid(BANDED); a.font = Font(color=COBALT, size=9)
    a.alignment = Alignment(horizontal="center", vertical="center")
    formulas = [
        f"=COUNTA(I{fila0}:I{ult})",
        f'=COUNTIF(I{fila0}:I{ult},"Aprobado")',
        f'=COUNTIF(I{fila0}:I{ult},"Fallido")',
        f'=COUNTIF(I{fila0}:I{ult},"Pendiente")',
        f'=COUNTIF(I{fila0}:I{ult},"En Progreso")',
        f'=COUNTIF(I{fila0}:I{ult},"Bloqueado")',
    ]
    for i, fml in enumerate(formulas):
        col = 9 + i
        cell = ws.cell(fv, col, fml)
        cell.fill = solid(GRAYBL); cell.font = Font(color=NAVY, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fv].height = 22

    ws.freeze_panes = "A5"
    return ult, fila0


def construir_dashboard(wb, meta, ult, fila0):
    ws = wb.create_sheet(DASH)
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 4), ("B", 20), ("C", 16), ("D", 16), ("E", 16), ("F", 4), ("G", 22), ("H", 12), ("I", 8)):
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:I1")
    t = ws["A1"]; t.value = "    DELTA TEST  ·  QA Dashboard"
    t.fill = solid(NAVY); t.font = Font(color=WHITE, bold=True, size=18)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells("A3:I3")
    ws["A3"].value = f"    Ambiente: {meta.get('ambiente','PRUEBAS')}  ·  Generado: {meta.get('fecha','')}"
    ws["A3"].font = Font(color=TXTHDR, size=10)

    rng = f"'{CHECK}'!$I${fila0}:$I$500"
    drng = f"'{CHECK}'!$D${fila0}:$D$500"
    cnt = lambda v: f'COUNTIF({rng},"{v}")'
    total_expr = "+".join(cnt(v) for v in ["Aprobado", "Fallido", "Pendiente", "En Progreso", "Bloqueado"])

    # Tarjetas de métricas
    tarjetas = [("B5", "Total Casos", f"={total_expr}"),
                ("C5", "Aprobados", f"={cnt('Aprobado')}"),
                ("D5", "Fallidos", f"={cnt('Fallido')}"),
                ("E5", "Pendientes", f"={cnt('Pendiente')}"),
                ("G5", "En Progreso", f"={cnt('En Progreso')}"),
                ("H5", "Bloqueados", f"={cnt('Bloqueado')}")]
    for celda, lbl, fml in tarjetas:
        col = celda[0]; r = int(celda[1:])
        h = ws[f"{col}{r}"]; h.value = lbl
        h.fill = solid(COBALT); h.font = Font(color=WHITE, bold=True, size=9)
        h.alignment = Alignment(horizontal="center")
        v = ws[f"{col}{r+1}"]; v.value = fml
        v.fill = solid(GRAYBL); v.font = Font(color=NAVY, bold=True, size=14)
        v.alignment = Alignment(horizontal="center")

    # Porcentajes
    pares = [("B8", "% Éxito", f"=IFERROR({cnt('Aprobado')}/({total_expr}),0)"),
             ("C8", "% Fallos", f"=IFERROR({cnt('Fallido')}/({total_expr}),0)"),
             ("D8", "% Avance", f"=IFERROR(({cnt('Aprobado')}+{cnt('Fallido')})/({total_expr}),0)")]
    for celda, lbl, fml in pares:
        col = celda[0]; r = int(celda[1:])
        ws[f"{col}{r}"].value = lbl
        ws[f"{col}{r}"].font = Font(bold=True, size=9, color=TXTHDR)
        pc = ws[f"{col}{r+1}"]; pc.value = fml; pc.number_format = "0.0%"
        pc.font = Font(color=NAVY, bold=True, size=12); pc.alignment = Alignment(horizontal="center")

    # Tabla Estado→Casos (para el gráfico de pastel)
    ws["B11"] = "Estado"; ws["C11"] = "Casos"
    estados = ["Aprobado", "Fallido", "Pendiente", "En Progreso", "Bloqueado"]
    for i, e in enumerate(estados):
        ws.cell(12 + i, 2, e)
        ws.cell(12 + i, 3, f"={cnt(e)}")
    for cc in ("B11", "C11"):
        ws[cc].font = Font(bold=True, color=WHITE, size=9); ws[cc].fill = solid(COBALT)
        ws[cc].alignment = Alignment(horizontal="center")

    # Tabla Módulo→Casos (para el gráfico de barras)
    ws["G11"] = "Módulo"; ws["H11"] = "Casos"
    modulos = ["HISTORIAS CLINICAS", "AGENDAMIENTO", "FACTURACIÓN", "LABORATORIO",
               "AUTENTICACIÓN", "CONSENTIMIENTOS", "ORDENAMIENTOS", "REIMPRESION", "NOTA"]
    for i, m in enumerate(modulos):
        ws.cell(12 + i, 7, m)
        ws.cell(12 + i, 8, f'=COUNTIF({drng},"{m}")')
    for cc in ("G11", "H11"):
        ws[cc].font = Font(bold=True, color=WHITE, size=9); ws[cc].fill = solid(COBALT)
        ws[cc].alignment = Alignment(horizontal="center")

    # Gráfico pastel (Estado)
    pie = PieChart(); pie.title = "Por Estado"
    data = Reference(ws, min_col=3, min_row=11, max_row=16)
    cats = Reference(ws, min_col=2, min_row=12, max_row=16)
    pie.add_data(data, titles_from_data=True); pie.set_categories(cats)
    pie.height = 7; pie.width = 11
    ws.add_chart(pie, "B19")

    # Gráfico de barras (Módulo)
    bar = BarChart(); bar.title = "Casos por Módulo"; bar.type = "col"; bar.legend = None
    bdata = Reference(ws, min_col=8, min_row=11, max_row=20)
    bcats = Reference(ws, min_col=7, min_row=12, max_row=20)
    bar.add_data(bdata, titles_from_data=True); bar.set_categories(bcats)
    bar.height = 7; bar.width = 13
    ws.add_chart(bar, "G19")

    # Resumen del ciclo
    resumen = [
        ("Proyecto", "Sistema de Historia Clínica"),
        ("Cliente", meta.get("cliente", "")),
        ("Ambiente", meta.get("ambiente", "PRUEBAS")),
        ("Versión", meta.get("version", "1.0.0")),
        ("Período", f"{meta.get('fecha','')} – {meta.get('fecha','')}"),
        ("Tester", meta.get("tester", "QA Automatization")),
        ("Responsable Dev", meta.get("desarrollador", "")),
        ("Criterio aceptación", "≥ 90% casos Aprobados"),
    ]
    ws.merge_cells("B34:I34")
    ws["B34"].value = "Resumen del Ciclo de Pruebas"
    ws["B34"].fill = solid(NAVY); ws["B34"].font = Font(color=WHITE, bold=True, size=11)
    for i, (k, v) in enumerate(resumen):
        r = 35 + i
        ws.cell(r, 2, k).font = Font(bold=True, color=TXTHDR, size=10)
        ws.merge_cells(f"C{r}:I{r}")
        ws.cell(r, 3, v).font = Font(color=TXT, size=10)


def main():
    json_path, xlsx_path = sys.argv[1], sys.argv[2]
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    meta = data["meta"]
    filas = data["filas"]

    wb = openpyxl.Workbook()
    ult, fila0 = construir_checklist(wb, meta, filas)
    construir_dashboard(wb, meta, ult, fila0)

    try:
        wb.save(xlsx_path)
    except PermissionError:
        base, ext = os.path.splitext(xlsx_path)
        xlsx_path = f"{base}_{datetime.now().strftime('%H%M%S')}{ext}"
        wb.save(xlsx_path)
    print(xlsx_path)


if __name__ == "__main__":
    main()
